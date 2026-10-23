from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Camion_Investor_Brief.docx"

BLUE = RGBColor(46, 116, 181)
NAVY = RGBColor(31, 77, 120)
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(84, 84, 84)
LIGHT_GRAY = "F4F6F9"
VERY_LIGHT_BLUE = "EAF2FB"


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), "9360")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_grid = tbl.tblGrid
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)


def style_paragraph(paragraph, *, size=11, color=BLACK, bold=False, italic=False, align=None, before=0, after=8, line_spacing=1.333):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line_spacing
    if align is not None:
        paragraph.alignment = align
    for run in paragraph.runs:
        set_run_font(run, size=size, color=color, bold=bold, italic=italic)


def add_text(paragraph, text, *, size=11, color=BLACK, bold=False, italic=False):
    run = paragraph.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return run


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    if level > 0:
        p.paragraph_format.left_indent = Inches(0.25 * level)
    add_text(p, text)
    style_paragraph(p, before=0, after=4, line_spacing=1.2)
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    add_text(p, text, size={1: 16, 2: 13, 3: 12}[level], color={1: BLUE, 2: BLUE, 3: NAVY}[level], bold=True)
    fmt = p.paragraph_format
    if level == 1:
        fmt.space_before = Pt(18)
        fmt.space_after = Pt(10)
    elif level == 2:
        fmt.space_before = Pt(12)
        fmt.space_after = Pt(6)
    else:
        fmt.space_before = Pt(8)
        fmt.space_after = Pt(4)
    fmt.line_spacing = 1.15
    return p


def add_callout(doc, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, VERY_LIGHT_BLUE)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_run_font(run, size=11, color=NAVY, bold=True)
    return table


def add_label_value_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_geometry(table, [1650, 7710])
    for i, (label, value) in enumerate(rows):
        left = table.cell(i, 0)
        right = table.cell(i, 1)
        set_cell_shading(left, LIGHT_GRAY)
        for cell in (left, right):
            cell.vertical_alignment = 1
        lp = left.paragraphs[0]
        rp = right.paragraphs[0]
        lp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        rp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        lp.paragraph_format.space_before = Pt(2)
        lp.paragraph_format.space_after = Pt(2)
        rp.paragraph_format.space_before = Pt(2)
        rp.paragraph_format.space_after = Pt(2)
        lp.paragraph_format.line_spacing = 1.15
        rp.paragraph_format.line_spacing = 1.15
        add_text(lp, label, size=10.5, color=GRAY, bold=True)
        add_text(rp, value, size=10.5, color=BLACK)
    return table


def add_matrix_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    set_table_geometry(table, widths)
    hdr = table.rows[0].cells
    for cell, header in zip(hdr, headers):
        set_cell_shading(cell, LIGHT_GRAY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.1
        add_text(p, header, size=10.5, color=NAVY, bold=True)
    for row_idx, row in enumerate(rows, start=1):
        for col_idx, value in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_text(p, value, size=10.5, color=BLACK)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    set_run_font(run, name="Calibri", size=9, color=GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run_el = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")
    rpr.append(rfonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")
    rpr.append(sz)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "545454")
    rpr.append(color)
    run_el.append(rpr)
    fld.append(run_el)
    paragraph._p.append(fld)


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.4)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, NAVY, 8, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15


def build_document():
    doc = Document()
    configure_styles(doc)

    section = doc.sections[0]
    configure_section(section)
    section.different_first_page_header_footer = True

    # Default running header/footer
    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_p.paragraph_format.space_before = Pt(0)
    header_p.paragraph_format.space_after = Pt(0)
    header_p.paragraph_format.line_spacing = 1
    add_text(header_p, "Camion Investor Brief", size=9, color=GRAY, bold=True)

    footer_p = section.footer.paragraphs[0]
    footer_p.paragraph_format.space_before = Pt(0)
    footer_p.paragraph_format.space_after = Pt(0)
    footer_p.paragraph_format.line_spacing = 1
    add_page_number(footer_p)

    # First page title block
    for _ in range(3):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(6)
    kicker.paragraph_format.line_spacing = 1
    add_text(kicker, "INVESTOR BRIEF", size=10, color=GRAY, bold=True)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    title.paragraph_format.line_spacing = 1
    add_text(title, "Camion", size=28, color=BLACK, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(10)
    subtitle.paragraph_format.line_spacing = 1.1
    add_text(
        subtitle,
        "A commerce platform that combines shopping, creator-led growth, and operational visibility in one mobile experience.",
        size=13,
        color=GRAY,
    )

    metadata = add_label_value_table(
        doc,
        [
            ("Prepared for", "Prospective investors and strategic partners"),
            ("Prepared by", "Camion product team"),
            ("Document type", "Investor overview"),
            ("Date", date(2026, 6, 7).strftime("%B %-d, %Y")),
        ],
    )
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()
    add_callout(
        doc,
        "Investment thesis: Camion is positioned as a hybrid commerce and growth engine. It does not only support online shopping; it also builds repeatable acquisition through affiliate and influencer monetization, multilingual commerce, and post-purchase engagement.",
    )

    doc.add_page_break()

    # Page 2 onward
    add_heading(doc, "Executive Summary", 1)
    p = doc.add_paragraph()
    add_text(
        p,
        "Camion is a cross-platform Flutter commerce application designed for modern retail workflows. The product brings together product discovery, cart and checkout, order management, notifications, wallet experiences, localization, and support tools in a single mobile app. From an investor perspective, the key story is not only software quality; it is that the app is structured to turn each customer into a higher-value repeat buyer and each creator into a measurable acquisition channel.",
    )
    style_paragraph(p)

    p = doc.add_paragraph()
    add_text(
        p,
        "The result is a commerce platform that can scale beyond a traditional storefront. It is built for markets where language, currency, delivery visibility, and social selling matter as much as the catalog itself.",
    )
    style_paragraph(p)

    add_heading(doc, "Purpose Of The Project", 1)
    p = doc.add_paragraph()
    add_text(
        p,
        "The purpose of Camion is to simplify the full customer journey and reduce friction at every step of the buying experience.",
    )
    style_paragraph(p, after=4)
    add_bullet(doc, "Make discovery and purchase easier through a mobile-first shopping experience.")
    add_bullet(doc, "Support multilingual and multi-currency commerce so the app can serve multiple markets with minimal product changes.")
    add_bullet(doc, "Improve trust after purchase with order tracking, notifications, wallet visibility, and support flows.")
    add_bullet(doc, "Create a built-in growth channel through affiliate and influencer dashboards, coupons, and commission tracking.")

    add_heading(doc, "Technology Used", 1)
    p = doc.add_paragraph()
    add_text(
        p,
        "The codebase reflects a modern Flutter architecture with a strong focus on modularity, performance, and configurable deployment.",
    )
    style_paragraph(p, after=6)

    add_matrix_table(
        doc,
        ["Layer", "Technologies In The Repo", "Why It Matters"],
        [
            (
                "Mobile app",
                "Flutter, Dart, flutter_bloc, go_router, flutter_screenutil, flutter_animate",
                "Cross-platform delivery, structured state management, and a flexible navigation model for a large app surface.",
            ),
            (
                "Data and networking",
                "Dio, REST integrations, GraphQL queries, environment flavors, model layer separation",
                "Supports maintainable API integration and future backend evolution without rewriting the client.",
            ),
            (
                "Identity and security",
                "Flutter secure storage, shared_preferences, Google/Facebook/Apple sign-in",
                "Enables secure sessions and low-friction login for conversion-focused mobile commerce.",
            ),
            (
                "Engagement and retention",
                "Firebase Messaging, analytics, crash reporting, local notifications, app links, in-app web views",
                "Improves re-engagement, conversion tracking, and operational reliability.",
            ),
            (
                "Commerce and content",
                "Cached images, cache management, video playback, photo viewing, carousel assets, wishlist support",
                "Creates a richer shopping experience that is ready for visual merchandising and content-driven selling.",
            ),
            (
                "Localization and market readiness",
                "intl, flutter_localizations, country-state-city utilities, currency handling",
                "Makes the product suitable for regional expansion and localized customer journeys.",
            ),
        ],
        [1400, 3300, 4660],
    )

    add_heading(doc, "Unique Business Model", 1)
    p = doc.add_paragraph()
    add_text(
        p,
        "Based on the implemented feature set, Camion can be positioned as a hybrid commerce platform with multiple monetization levers rather than a single-store app.",
    )
    style_paragraph(p, after=6)

    add_matrix_table(
        doc,
        ["Monetization Lever", "How It Works In The Product", "Investor Value"],
        [
            (
                "Commerce revenue",
                "The core shopping flow supports product browsing, carting, checkout, payment, and order management.",
                "Keeps the business grounded in direct transaction value.",
            ),
            (
                "Affiliate and influencer commissions",
                "Creators can apply, get approved, create coupons, and track wallet balances and coupon performance.",
                "Turns user acquisition into a measurable, performance-based channel.",
            ),
            (
                "Repeat purchase lift",
                "Wish lists, notifications, tracking, and support tools encourage customers to return and reorder.",
                "Improves lifetime value and reduces the dependence on paid acquisition.",
            ),
            (
                "Market expansion",
                "Flavor-based deployment, env-driven URLs, and localization/multi-currency support reduce friction when entering new regions.",
                "Makes expansion operationally lighter and more repeatable.",
            ),
        ],
        [1800, 3750, 3810],
    )

    p = doc.add_paragraph()
    add_text(
        p,
        "In simple terms, Camion is designed to earn from selling products and from enabling creators to sell on its behalf. That hybrid model is unusual because acquisition, conversion, and retention are all engineered inside the same app experience.",
    )
    style_paragraph(p)

    add_heading(doc, "Unique Selling Proposition", 1)
    add_bullet(doc, "The product is not only an e-commerce app; it is a commerce growth system with affiliate and influencer monetization built in.")
    add_bullet(doc, "It is multilingual and multi-currency ready, which is critical for regional scale and cross-border usage.")
    add_bullet(doc, "It provides the full post-purchase loop, including tracking, notifications, wallet visibility, and support, which improves trust and repeat ordering.")
    add_bullet(doc, "It is built on Flutter with modular architecture, which helps the team ship faster across iOS and Android from one codebase.")

    add_heading(doc, "Why This Is Investable", 1)
    add_bullet(doc, "Scalable product foundation: the app is already organized into distinct features for commerce, profile, wallet, notifications, search, and creator programs.")
    add_bullet(doc, "Operational flexibility: environment-based configuration and flavor support make it easier to separate development and production workflows.")
    add_bullet(doc, "Growth efficiency: the affiliate/influencer layer creates a performance-based acquisition model rather than depending only on paid media.")
    add_bullet(doc, "Expansion readiness: localization, currency support, and country-aware flows make the platform more adaptable to new markets.")

    add_heading(doc, "Closing Positioning", 1)
    p = doc.add_paragraph()
    add_text(
        p,
        "Camion should be presented to investors as a modern commerce platform with a differentiated growth engine. Its strength is the combination of product depth and monetization design: a polished shopping experience, measurable creator-led acquisition, and the technical flexibility needed to scale across markets.",
    )
    style_paragraph(p)

    add_heading(doc, "Strategic Narrative Draft", 1)
    p = doc.add_paragraph()
    add_text(
        p,
        "The following language is adapted from your founder narrative and positioned as a stronger strategic appendix. It frames Camion as infrastructure for commerce, not just a consumer app.",
    )
    style_paragraph(p, after=6)
    add_callout(
        doc,
        "This section is intentionally more ambitious than the core repo-based summary. Before external use, every operational and legal claim should be reviewed and verified.",
    )

    add_heading(doc, "Core Thesis", 2)
    p = doc.add_paragraph()
    add_text(
        p,
        "Camion is not just an e-commerce company. It is building the infrastructure layer of global commerce by turning fragmented, operations-heavy trade into a software-orchestrated system.",
    )
    style_paragraph(p)

    add_heading(doc, "What The Model Solves", 2)
    add_bullet(doc, "It reduces dependence on manual coordination by embedding workflow, visibility, and execution into software.")
    add_bullet(doc, "It addresses the inefficiency of fragmented supply, inflated intermediary layers, and inventory-heavy models.")
    add_bullet(doc, "It is designed to scale across markets without requiring a new operational stack for every geography.")

    add_heading(doc, "Why The Company Has A Technology Moat", 2)
    add_bullet(doc, "Software is the operating layer of the business, not a support function.")
    add_bullet(doc, "The system can coordinate procurement, logistics, margin logic, and fulfillment in a unified workflow.")
    add_bullet(doc, "As the platform scales, operational data and execution complexity strengthen the moat.")

    add_heading(doc, "Economic Logic", 2)
    add_bullet(doc, "Direct sourcing lowers cost of goods.")
    add_bullet(doc, "Consolidation improves logistics efficiency.")
    add_bullet(doc, "On-demand execution reduces inventory risk.")
    add_bullet(doc, "System-driven pricing and orchestration improve margin discipline.")

    add_heading(doc, "Why Now", 2)
    p = doc.add_paragraph()
    add_text(
        p,
        "The market is moving toward lower-cost, software-defined, cross-border commerce. Camion is positioned to ride that shift rather than react to it.",
    )
    style_paragraph(p)

    add_heading(doc, "Strategic Positioning Statement", 2)
    p = doc.add_paragraph()
    add_text(
        p,
        "We are not building a store. We are building the system that commerce runs on.",
    )
    style_paragraph(p, after=0)

    p = doc.add_paragraph()
    add_text(
        p,
        "If you want, the next step is to turn this into a branded investor pitch deck or a shorter one-page teaser for outreach.",
    )
    style_paragraph(p, after=0)

    return doc


def main():
    doc = build_document()
    doc.core_properties.title = "Camion Investor Brief"
    doc.core_properties.subject = "Investor-facing project overview"
    doc.core_properties.author = "OpenAI Codex"
    doc.core_properties.company = "Camion"
    doc.core_properties.comments = "Generated from repository analysis"
    doc.save(OUTPUT)
    print(f"Saved {OUTPUT}")


if __name__ == "__main__":
    main()
